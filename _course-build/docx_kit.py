"""Word-file helpers shared by the module builders (same look as the Module 1 files)."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RED = RGBColor(0xD7, 0x1F, 0x2D)
BLUE = RGBColor(0x3A, 0x56, 0xA6)
INK = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5A, 0x63, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Word requires tblPr children in schema order
ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
         "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def borders(table, color="B9C2D4"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    table._tbl.tblPr.append(b)


def fix(table, widths):
    """Fixed layout with explicit widths so Word, Google Docs and previews agree."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    return table


def row_height(row, inches):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def bottom_rule(p, color="D71F2D", size="12"):
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), size)
    bt.set(qn("w:color"), color)
    bt.set(qn("w:space"), "4")
    pb.append(bt)
    pPr.insert(0, pb)


class Kit:
    def __init__(self, logo):
        self.logo = logo
        self.doc = doc = Document()
        self.sec = sec = doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(0.9))
        st = doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(4)
        for name, size, color in (("Heading 1", 20, INK), ("Heading 2", 14, BLUE), ("Heading 3", 11.5, INK)):
            h = doc.styles[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(12 if name != "Heading 3" else 8)
            h.paragraph_format.space_after = Pt(4)

    def para(self, text="", bold=False, italic=False, muted=False, size=None, after=None):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if muted:
            r.font.color.rgb = MUTED
        if size:
            r.font.size = Pt(size)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def header(self, kicker, title):
        head = self.doc.add_table(rows=1, cols=2)
        fix(head, [5.6, 1.1])
        p = head.cell(0, 0).paragraphs[0]
        r = p.add_run(kicker)
        r.font.size, r.font.color.rgb, r.bold = Pt(9), MUTED, True
        r = head.cell(0, 0).add_paragraph().add_run(title)
        r.font.size, r.bold, r.font.color.rgb = Pt(20), True, INK
        lp = head.cell(0, 1).paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lp.add_run().add_picture(self.logo, width=Inches(0.8))
        bottom_rule(self.para("Chart Academy · in partnership with IRSJA", muted=True, size=9, after=10))

    def name_date(self):
        info = self.doc.add_table(rows=1, cols=3)
        borders(info)
        for i, label in enumerate(("Name:", "Date:", "Cohort:")):
            r = info.cell(0, i).paragraphs[0].add_run(label)
            r.bold = True
            r.font.size = Pt(10)
        row_height(info.rows[0], 0.4)
        fix(info, [2.9, 1.9, 1.9])
        self.para("", after=2)

    def bullets(self, lines):
        for lead, body in lines:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(lead).bold = True
            p.add_run(body)

    def callout(self, title, body, fill, title_hex):
        t = self.doc.add_table(rows=1, cols=1)
        fix(t, [6.7])
        c = t.cell(0, 0)
        shade(c, fill)
        p = c.paragraphs[0]
        r = p.add_run(title + "  ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(title_hex)
        p.add_run(body)
        self.para("", after=2)

    def answer_box(self, prompt, height=1.0, hint=None):
        p = self.para(prompt, bold=True, after=2)
        if hint:
            r = p.add_run("  " + hint)
            r.bold, r.italic = False, True
            r.font.color.rgb = MUTED
            r.font.size = Pt(9.5)
        t = self.doc.add_table(rows=1, cols=1)
        borders(t)
        fix(t, [6.7])
        shade(t.cell(0, 0), "F6F8FC")
        row_height(t.rows[0], height)
        self.para("", after=2)

    def grid(self, headers, widths, rows, height=0.5, first_col=None, header_fill="3A56A6"):
        t = self.doc.add_table(rows=1 + rows, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t)
        for i, h in enumerate(headers):
            c = t.cell(0, i)
            shade(c, header_fill)
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.color.rgb = WHITE
            r.font.size = Pt(10)
        for ri in range(1, rows + 1):
            row_height(t.rows[ri], height)
            if first_col and ri - 1 < len(first_col) and first_col[ri - 1]:
                r = t.cell(ri, 0).paragraphs[0].add_run(first_col[ri - 1])
                r.bold = True
                r.font.size = Pt(10)
        fix(t, widths)
        self.para("", after=2)
        return t

    def cell_text(self, table, row, col, text):
        table.cell(row, col).paragraphs[0].add_run(text).font.size = Pt(10)

    def role_picker(self, names):
        half = (len(names) + 1) // 2
        t = self.grid(["", "Role", "", "Role"], [0.4, 2.95, 0.4, 2.95], half, height=0.3)
        for i, n in enumerate(names):
            self.cell_text(t, 1 + i % half, 1 if i < half else 3, n)

    def rubric(self, rows):
        doc = self.doc
        doc.add_page_break()
        doc.add_heading("Rubric · 20 points", level=2)
        rub = doc.add_table(rows=len(rows) + 2, cols=4)
        borders(rub)
        w = [1.6, 1.7, 1.7, 1.7]
        for i, h in enumerate(["", "Full credit · 5", "Partial · 3", "Not yet · 1"]):
            c = rub.cell(0, i)
            shade(c, "3A56A6")
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(10), WHITE
        for ri, row in enumerate(rows, start=1):
            for ci, text in enumerate(row):
                r = rub.cell(ri, ci).paragraphs[0].add_run(text)
                r.font.size = Pt(9.5)
                r.bold = ci == 0
        fix(rub, w)
        last = rub.rows[len(rows) + 1]
        merged = last.cells[0].merge(last.cells[2])
        merged.paragraphs[0].add_run("Total").bold = True
        shade(merged, "F6F8FC")
        shade(last.cells[3], "F6F8FC")
        last.cells[3].paragraphs[0].add_run("      / 20").bold = True
        self.para("")
        self.para("Instructor feedback", bold=True, after=2)
        t = doc.add_table(rows=1, cols=1)
        borders(t)
        fix(t, [6.7])
        row_height(t.rows[0], 1.4)

    def footer(self, text):
        fp = self.sec.footer.paragraphs[0]
        fp.text = text
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = MUTED

    def save(self, path):
        for z in self.doc.settings.element.iter(qn("w:zoom")):
            z.set(qn("w:percent"), "100")
        for tblPr in self.doc.element.body.iter(qn("w:tblPr")):
            kids = list(tblPr)
            for k in kids:
                tblPr.remove(k)
            kids.sort(key=lambda e: ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDER else 99)
            for k in kids:
                tblPr.append(k)
        self.doc.save(path)
