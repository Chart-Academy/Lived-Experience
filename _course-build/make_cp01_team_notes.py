"""Build the Module 1 team notes as a shared Word file."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = sys.argv[1]
LOGO = sys.argv[2]
RED = RGBColor(0xD7, 0x1F, 0x2D)
BLUE = RGBColor(0x3A, 0x56, 0xA6)
INK = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5A, 0x63, 0x78)
GREEN = RGBColor(0x3F, 0x6B, 0x50)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(0.9))

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(4)
for name, size, color in (("Heading 1", 20, INK), ("Heading 2", 14, BLUE), ("Heading 3", 11.5, INK)):
    h = doc.styles[name]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = color
    h.paragraph_format.space_before = Pt(12 if name != "Heading 3" else 8)
    h.paragraph_format.space_after = Pt(4)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def borders(table, color="B9C2D4"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def fix(table, widths):
    """Fixed layout with explicit widths so Word, Google Docs and previews agree."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid_el = table._tbl.tblGrid
    for gc, w in zip(grid_el.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    return table


def row_height(row, inches):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def bottom_rule(p, color="D71F2D", size="12"):
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), size)
    bt.set(qn("w:color"), color)
    bt.set(qn("w:space"), "4")
    pb.append(bt)
    pPr.insert(0, pb)


def para(text="", bold=False, italic=False, color=None, size=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def answer_box(prompt, height=1.0, hint=None):
    """A question followed by a shaded box the learner types into."""
    p = para(prompt, bold=True, after=2)
    if hint:
        r = p.add_run("  " + hint)
        r.bold = False
        r.italic = True
        r.font.color.rgb = MUTED
        r.font.size = Pt(9.5)
    t = doc.add_table(rows=1, cols=1)
    borders(t)
    t.columns[0].width = Inches(6.7)
    fix(t, [6.7])
    c = t.cell(0, 0)
    shade(c, "F6F8FC")
    row_height(t.rows[0], height)
    c.paragraphs[0].add_run("")
    para("", after=2)


def grid(headers, widths, rows, height=0.5, header_fill="3A56A6", first_col=None):
    t = doc.add_table(rows=1 + rows, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.cell(0, i)
        c.width = Inches(w)
        shade(c, header_fill)
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
    for ri in range(1, rows + 1):
        row_height(t.rows[ri], height)
        if first_col:
            r = t.cell(ri, 0).paragraphs[0].add_run(first_col[ri - 1])
            r.bold = True
            r.font.size = Pt(10)
    fix(t, widths)
    para("", after=2)
    return t


def callout(title, body, fill, title_color):
    t = doc.add_table(rows=1, cols=1)
    fix(t, [6.7])
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = title_color
    p.add_run(body)
    para("", after=2)


# Header
head = doc.add_table(rows=1, cols=2)
fix(head, [5.6, 1.1])
p = head.cell(0, 0).paragraphs[0]
r = p.add_run("CERTIFIED COMMUNITY PRACTITIONER · MODULE 01 · TEAM NOTES")
r.font.size, r.font.color.rgb, r.bold = Pt(9), MUTED, True
p2 = head.cell(0, 0).add_paragraph()
r = p2.add_run("Team Notes: Denise's Systems Map")
r.font.size, r.bold, r.font.color.rgb = Pt(20), True, INK
lp = head.cell(0, 1).paragraphs[0]
lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
lp.add_run().add_picture(LOGO, width=Inches(0.8))
rule = para("Chart Academy · in partnership with IRSJA", color=MUTED, size=9, after=10)
bottom_rule(rule)

doc.add_heading("How to use these notes", level=2)
for a, b in (
    ("One note-taker: ", "The team lead (or a volunteer) types the team's notes here during the huddle."),
    ("Or work together: ", "Upload this file to Google Drive, open it with Google Docs, and share it with your team so everyone can type at once."),
    ("Before you leave: ", "Everyone on the team uses these notes to finish Part 2 of their own Module 1 assignment. The note-taker uploads these notes to Canvas with the team's names."),
):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(a).bold = True
    p.add_run(b)
callout("Remember:", "Denise is fictional. Never include real client names or details.", "FFF8E1", RGBColor(0x8A, 0x69, 0x00))

doc.add_heading("Our team", level=2)
grid(["Name", "Role"], [3.35, 3.35], 6, height=0.35,
     first_col=None)
t = doc.tables[-1]
for i, role in enumerate(["Housing navigator", "Benefits outreach worker", "Reentry navigator",
                          "Community health worker", "Workforce specialist", "Team lead"], start=1):
    run = t.cell(i, 1).paragraphs[0].add_run(role)
    run.font.size = Pt(10)

doc.add_heading("The case", level=2)
para("Denise, 34, is staying in a family shelter with her two children, ages 4 and 9. She came home from jail "
     "10 months ago. She lost her retail job when the store closed, her SNAP benefits are about to run out, "
     "and her youngest child has asthma.")
para("“Every office wants something from me before they'll help. I just need one person to help me figure out "
     "where to start.”", italic=True)

doc.add_heading("Step 1 · What each role knows", level=2)
para("Each role shares what is on their card. The note-taker fills in that row.", italic=True, color=MUTED)
grid(["Role", "Denise's needs in this system", "What this agency can offer", "What this agency needs from another"],
     [1.45, 1.75, 1.75, 1.75], 6, height=0.75,
     first_col=["Housing", "Benefits", "Reentry / justice", "Health", "Employment", "Other (ID, child care)"])

doc.add_heading("Step 2 · Connections and conflicts", level=2)
para("Where does one system depend on or block another? Mark the places where Denise is stuck.",
     italic=True, color=MUTED)
grid(["This role…", "needs this…", "from this role", "Stuck? (Y/N)"], [1.6, 2.5, 1.6, 1.0], 5, height=0.5)
para("Optional: take a photo of your team's paper map and add it at the end of this document.",
     italic=True, color=MUTED, size=9.5)

doc.add_heading("Step 3 · Our plan", level=2)
answer_box("Denise's first step, and why it comes first:", 0.9)
answer_box("Lead practitioner (Denise's main contact), and why:", 0.6)
answer_box("Warm handoff: who calls whom, and what do they say?", 1.1, hint="(Write it the way you'd say it on the phone.)")

doc.add_heading("Presentation · 2 minutes", level=2)
para("Each person says one sentence in role.", italic=True, color=MUTED)
grid(["Role", "My sentence"], [1.8, 4.9], 6, height=0.45,
     first_col=["Housing", "Benefits", "Reentry", "Health", "Employment", "Team lead"])

fp = sec.footer.paragraphs[0]
fp.text = "Chart Academy Workforce Credentialing · Certified Community Practitioner · Module 01 Team Notes"
fp.runs[0].font.size = Pt(8.5)
fp.runs[0].font.color.rgb = MUTED

for z in doc.settings.element.iter(qn("w:zoom")):
    z.set(qn("w:percent"), "100")
ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
         "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
for tblPr in doc.element.body.iter(qn("w:tblPr")):
    kids = list(tblPr)
    for k in kids:
        tblPr.remove(k)
    kids.sort(key=lambda e: ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDER else 99)
    for k in kids:
        tblPr.append(k)

doc.save(OUT)
print("saved", OUT)
