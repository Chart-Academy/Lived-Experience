"""Build the Module 1 Community Practitioner assignment as a fillable Word file."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = sys.argv[1]
LOGO = sys.argv[2]
RED = RGBColor(0xD7, 0x1F, 0x2D)
BLUE = RGBColor(0x3A, 0x56, 0xA6)
INK = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5A, 0x63, 0x78)
GREEN = RGBColor(0x3F, 0x6B, 0x50)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(0.9))

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(4)
for name, size, color in (("Heading 1", 20, INK), ("Heading 2", 14, BLUE), ("Heading 3", 11.5, INK)):
    h = doc.styles[name]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = color
    h.paragraph_format.space_before = Pt(12 if name != "Heading 3" else 8)
    h.paragraph_format.space_after = Pt(4)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def borders(table, color="B9C2D4"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def fix(table, widths):
    """Fixed layout with explicit widths so Word, Google Docs and previews agree."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid_el = table._tbl.tblGrid
    for gc, w in zip(grid_el.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    return table


def row_height(row, inches):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def bottom_rule(p, color="D71F2D", size="12"):
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), size)
    bt.set(qn("w:color"), color)
    bt.set(qn("w:space"), "4")
    pb.append(bt)
    pPr.insert(0, pb)


def para(text="", bold=False, italic=False, color=None, size=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def answer_box(prompt, height=1.0, hint=None):
    """A question followed by a shaded box the learner types into."""
    p = para(prompt, bold=True, after=2)
    if hint:
        r = p.add_run("  " + hint)
        r.bold = False
        r.italic = True
        r.font.color.rgb = MUTED
        r.font.size = Pt(9.5)
    t = doc.add_table(rows=1, cols=1)
    borders(t)
    t.columns[0].width = Inches(6.7)
    fix(t, [6.7])
    c = t.cell(0, 0)
    shade(c, "F6F8FC")
    row_height(t.rows[0], height)
    c.paragraphs[0].add_run("")
    para("", after=2)


def grid(headers, widths, rows, height=0.5, header_fill="3A56A6", first_col=None):
    t = doc.add_table(rows=1 + rows, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.cell(0, i)
        c.width = Inches(w)
        shade(c, header_fill)
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
    for ri in range(1, rows + 1):
        row_height(t.rows[ri], height)
        if first_col:
            r = t.cell(ri, 0).paragraphs[0].add_run(first_col[ri - 1])
            r.bold = True
            r.font.size = Pt(10)
    fix(t, widths)
    para("", after=2)
    return t


def callout(title, body, fill, title_color):
    t = doc.add_table(rows=1, cols=1)
    fix(t, [6.7])
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = title_color
    p.add_run(body)
    para("", after=2)


# Header
head = doc.add_table(rows=1, cols=2)
fix(head, [5.6, 1.1])
p = head.cell(0, 0).paragraphs[0]
r = p.add_run("CERTIFIED COMMUNITY PRACTITIONER · MODULE 01 · ASSIGNMENT")
r.font.size, r.font.color.rgb, r.bold = Pt(9), MUTED, True
p2 = head.cell(0, 0).add_paragraph()
r = p2.add_run("The Human Services Practitioner Role")
r.font.size, r.bold, r.font.color.rgb = Pt(20), True, INK
lp = head.cell(0, 1).paragraphs[0]
lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
lp.add_run().add_picture(LOGO, width=Inches(0.8))
rule = para("Chart Academy · in partnership with IRSJA", color=MUTED, size=9, after=10)
bottom_rule(rule)

info = doc.add_table(rows=1, cols=3)
borders(info)
for i, label in enumerate(("Name:", "Date:", "Cohort:")):
    c = info.cell(0, i)
    c.width = Inches(2.23)
    r = c.paragraphs[0].add_run(label)
    r.bold = True
    r.font.size = Pt(10)
row_height(info.rows[0], 0.4)
fix(info, [2.9, 1.9, 1.9])
para("", after=2)

doc.add_heading("About this assignment", level=2)
para("In class, you played a practitioner on Denise's team and helped build her systems map. "
     "This assignment captures what you learned, in your own words. You complete it in class. There is no homework.")
for line in (
    ("How to complete it: ", "Download this file, type your answers into the shaded boxes, and save it. Use your team notes for Part 2."),
    ("How to submit: ", "Upload your saved file to the Module 1 assignment in Canvas. "
                        "Working on paper? Print it, fill it in, and upload a clear photo of each page."),
    ("Due: ", "By the end of class today. Upload it before you leave."),
    ("Points: ", "20 (see the rubric on the last page)."),
):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(line[0])
    r.bold = True
    p.add_run(line[1])
callout("Remember:", "Denise is fictional. Never include real client names or details.", "FFF8E1", RGBColor(0x8A, 0x69, 0x00))

# Part 1
doc.add_heading("Part 1 · Your role on the team", level=2)
para("Answer these from the point of view of the role you played in class.", italic=True, color=MUTED)
para("1. Which role did you play? Mark one with an X.", bold=True, after=2)
roles = grid(["", "Role", "", "Role"], [0.4, 2.95, 0.4, 2.95], 3, height=0.3)
names = ["Housing navigator", "Benefits outreach worker", "Reentry navigator",
         "Community health worker", "Workforce specialist", "Team lead"]
for i, n in enumerate(names):
    roles.cell(1 + i % 3, 1 if i < 3 else 3).paragraphs[0].add_run(n).font.size = Pt(10)
answer_box("2. What did your role know about Denise that the rest of the team didn't?", 0.9)
answer_box("3. What did your agency need from another agency before it could help Denise?", 0.8)
answer_box("4. Name one thing your role should NOT do. Who would you refer Denise to instead?", 0.8,
           hint="(Think about the limits of peer, practitioner, and clinical roles.)")

# Part 2
doc.add_heading("Part 2 · Your team's systems map", level=2)
para("Use your team's map from class. For each system, write Denise's needs and an agency or program that can help.",
     italic=True, color=MUTED)
grid(["System", "Denise's needs", "Agency or program that can help"], [1.5, 2.6, 2.6], 6, height=0.55,
     first_col=["Housing", "Benefits", "Reentry / justice", "Health", "Employment", "Other (ID, child care)"])
answer_box("5. Name two places where one system blocked or depended on another.", 0.9,
           hint="(Example: she can't do X until Y happens.)")
answer_box("6. What was your team's first step for Denise, and why does it come first?", 0.9)
answer_box("7. Write the warm handoff. Who is calling whom, and what do they say?", 1.1,
           hint="(2–4 sentences, written as if you are on the phone.)")
para("Optional: add a photo of your team's map at the end of this document.", italic=True, color=MUTED, size=9.5)

# Part 3
doc.add_heading("Part 3 · Lived experience reflection", level=2)
callout("In your own words.", "Write this part yourself. You choose what to share, and you never have to share "
        "more than you're comfortable with. Your strengths will go into your résumé during Job Search.",
        "EDF3EE", GREEN)
para("8. Name three strengths from your own life that you will bring to this work.", bold=True, after=2)
grid(["Strength", "Where it came from in my life", "What it looks like on the job"], [1.8, 2.45, 2.45], 3,
     height=0.75, first_col=["1.", "2.", "3."])
answer_box("9. Today you sat on the practitioner's side of the desk. Have you ever been in a situation like "
           "Denise's, or helped someone who was? What did that experience teach you that will make you a "
           "better practitioner?", 2.4, hint="(150–250 words.)")
answer_box("10. Finish this sentence: “A strength I bring to this work from my own life is … because …”", 0.8)

# Rubric
doc.add_page_break()
doc.add_heading("Rubric · 20 points", level=2)
rub = doc.add_table(rows=6, cols=4)
borders(rub)
w = [1.6, 1.7, 1.7, 1.7]
hdr = ["", "Full credit · 5", "Partial · 3", "Not yet · 1"]
rows = [
    ["Role and limits (Part 1)", "Explains what the role knew and needed, and names a clear limit and the right referral.",
     "Answers most questions. The limit or referral is unclear.", "Missing answers, or the role steps outside its limits."],
    ["Systems map (Part 2, table and question 5)", "Needs and a resource for every system, with two clear connections or conflicts.",
     "Most systems complete. One connection, or connections are vague.", "Few systems complete. No connections."],
    ["Plan and handoff (questions 6–7)", "First step is explained. The handoff is specific, respectful, and realistic.",
     "First step or handoff is present but not explained.", "Missing, or the handoff is a cold referral."],
    ["Lived experience reflection (Part 3)", "Three specific strengths, each tied to life and to the job. Reflection is thoughtful and in the learner's own words.",
     "Strengths are general, or the reflection is short.", "Missing, or not in the learner's own words."],
]
for i, h in enumerate(hdr):
    c = rub.cell(0, i)
    c.width = Inches(w[i])
    shade(c, "3A56A6")
    r = c.paragraphs[0].add_run(h)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10), RGBColor(0xFF, 0xFF, 0xFF)
for ri, row in enumerate(rows, start=1):
    for ci, text in enumerate(row):
        c = rub.cell(ri, ci)
        c.width = Inches(w[ci])
        r = c.paragraphs[0].add_run(text)
        r.font.size = Pt(9.5)
        r.bold = ci == 0
fix(rub, w)
last = rub.rows[5]
merged = last.cells[0].merge(last.cells[2])
merged.paragraphs[0].add_run("Total").bold = True
shade(merged, "F6F8FC")
shade(last.cells[3], "F6F8FC")
last.cells[3].paragraphs[0].add_run("      / 20").bold = True

para("")
para("Instructor feedback", bold=True, after=2)
t = doc.add_table(rows=1, cols=1)
borders(t)
fix(t, [6.7])
row_height(t.rows[0], 1.4)

# Footer
fp = sec.footer.paragraphs[0]
fp.text = "Chart Academy Workforce Credentialing · Certified Community Practitioner · Module 01 Assignment"
fp.runs[0].font.size = Pt(8.5)
fp.runs[0].font.color.rgb = MUTED

for z in doc.settings.element.iter(qn("w:zoom")):
    z.set(qn("w:percent"), "100")

# Word requires tblPr children in schema order
ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
         "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
for tblPr in doc.element.body.iter(qn("w:tblPr")):
    kids = list(tblPr)
    for k in kids:
        tblPr.remove(k)
    kids.sort(key=lambda e: ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDER else 99)
    for k in kids:
        tblPr.append(k)

doc.save(OUT)
print("saved", OUT)
